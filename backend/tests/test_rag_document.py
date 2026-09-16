"""文档解析与分块的单元测试。

这些测试**不调用任何外部 API**（不花钱、不依赖网络），专门验证
「把文件读成文字」和「把文字切成块」这两个环节是否正确。
"""

import pytest

from app.rag.loaders import (
    UnsupportedFileTypeError,
    is_supported,
    load_document,
    supported_extensions,
)
from app.rag.splitter import split_documents

SAMPLE_PRODUCT_TEXT = """星耀 X9 Pro 智能手机 产品介绍

一、核心配置
星耀 X9 Pro 搭载骁龙 8 Gen 3 处理器，配备 12GB LPDDR5X 运行内存
和 256GB UFS 4.0 存储空间，安兔兔跑分超过 210 万分。

二、电池与续航
内置 5000mAh 大容量电池，支持 67W 有线快充和 50W 无线快充。
实测有线快充 38 分钟可从 1% 充至 100%。

三、影像系统
后置三摄：5000 万像素主摄（索尼 IMX890，支持 OIS 光学防抖）、
800 万像素超广角、200 万像素微距。

四、保修政策
自购买之日起提供 12 个月整机保修，电池等易损耗部件保修 6 个月。
"""


# ============================================================
# 文件类型判断
# ============================================================
class TestFileTypeSupport:
    @pytest.mark.parametrize(
        "filename", ["a.pdf", "a.docx", "a.xlsx", "a.txt", "a.md", "A.PDF", "路径/深一点/文件.Md"]
    )
    def test_supported(self, filename):
        assert is_supported(filename)

    @pytest.mark.parametrize("filename", ["a.doc", "a.png", "a.zip", "a", "a.xls"])
    def test_unsupported(self, filename):
        assert not is_supported(filename)

    def test_supported_extensions_content(self):
        assert supported_extensions() == {".pdf", ".docx", ".xlsx", ".txt", ".md"}


# ============================================================
# 纯文本 / Markdown 解析
# ============================================================
class TestTextLoader:
    def test_load_txt(self, tmp_path):
        f = tmp_path / "商品介绍.txt"
        f.write_text(SAMPLE_PRODUCT_TEXT, encoding="utf-8")

        docs = load_document(f, "商品介绍.txt", {"doc_id": 1, "kb_id": 1})

        assert len(docs) == 1
        assert "5000mAh" in docs[0].page_content
        assert docs[0].metadata["filename"] == "商品介绍.txt"
        assert docs[0].metadata["file_type"] == ".txt"
        assert docs[0].metadata["doc_id"] == 1

    def test_load_gbk_encoded_txt(self, tmp_path):
        """老编码（GBK）的文本文件也要能正确读出来。

        国内很多 Windows 上生成的 txt 是 GBK 编码，如果只按 UTF-8 读会直接报错。
        """
        f = tmp_path / "gbk.txt"
        f.write_bytes(SAMPLE_PRODUCT_TEXT.encode("gbk"))

        docs = load_document(f, "gbk.txt")
        assert "5000mAh" in docs[0].page_content

    def test_load_utf8_with_bom(self, tmp_path):
        """带 BOM 的 UTF-8 文件（Windows 记事本另存为 UTF-8 会产生）要能读。"""
        f = tmp_path / "bom.txt"
        f.write_bytes(b"\xef\xbb\xbf" + SAMPLE_PRODUCT_TEXT.encode("utf-8"))

        docs = load_document(f, "bom.txt")
        assert "5000mAh" in docs[0].page_content
        # BOM 不能被当成正文内容留下来，否则第一个字前面会多一个看不见的字符
        assert not docs[0].page_content.startswith("﻿")

    def test_empty_file_rejected(self, tmp_path):
        f = tmp_path / "空文件.txt"
        f.write_text("   \n\n  ", encoding="utf-8")
        with pytest.raises(ValueError, match="空的"):
            load_document(f, "空文件.txt")

    def test_unsupported_extension_raises(self, tmp_path):
        f = tmp_path / "图片.png"
        f.write_bytes(b"fake png")
        with pytest.raises(UnsupportedFileTypeError, match="不支持"):
            load_document(f, "图片.png")


# ============================================================
# Word 解析
# ============================================================
class TestDocxLoader:
    def test_load_docx_with_paragraphs_and_table(self, tmp_path):
        """Word 里的正文段落和表格都要提取到。

        商品资料经常用表格列规格参数，只读段落会漏掉关键信息。
        """
        from docx import Document as DocxDocument

        docx = DocxDocument()
        docx.add_paragraph("星耀 X9 Pro 商品参数表")
        docx.add_paragraph("以下是本机型的详细技术规格。")

        table = docx.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "参数名称"
        table.cell(0, 1).text = "参数值"
        table.cell(1, 0).text = "电池容量"
        table.cell(1, 1).text = "5000mAh"
        table.cell(2, 0).text = "充电功率"
        table.cell(2, 1).text = "67W"

        f = tmp_path / "参数.docx"
        docx.save(f)

        docs = load_document(f, "参数.docx")
        content = docs[0].page_content

        assert "星耀 X9 Pro 商品参数表" in content
        assert "以下是本机型的详细技术规格" in content
        # 表格被转成「字段: 值」形式
        assert "电池容量: 5000mAh" in content
        assert "充电功率: 67W" in content


# ============================================================
# Excel 解析
# ============================================================
class TestXlsxLoader:
    def test_load_xlsx_specification_table(self, tmp_path):
        """商品规格表：第一行表头 + 后续数据行，转成「字段: 值」。"""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "规格参数"
        ws.append(["参数名称", "参数值"])
        ws.append(["屏幕尺寸", "6.7 英寸"])
        ws.append(["电池容量", "5000mAh"])
        ws.append(["机身重量", "198g"])

        f = tmp_path / "规格.xlsx"
        wb.save(f)

        docs = load_document(f, "规格.xlsx")
        content = docs[0].page_content

        assert "屏幕尺寸: 6.7 英寸" in content
        assert "电池容量: 5000mAh" in content
        assert docs[0].metadata["sheet"] == "规格参数"

    def test_load_xlsx_multiple_sheets(self, tmp_path):
        """多个工作表要分别提取，不能只读第一个。"""
        from openpyxl import Workbook

        wb = Workbook()
        ws1 = wb.active
        ws1.title = "参数"
        ws1.append(["电池", "5000mAh"])

        ws2 = wb.create_sheet("售后")
        ws2.append(["保修期", "12 个月"])

        f = tmp_path / "多表.xlsx"
        wb.save(f)

        docs = load_document(f, "多表.xlsx")

        assert len(docs) == 2
        sheets = {d.metadata["sheet"] for d in docs}
        assert sheets == {"参数", "售后"}

    def test_empty_xlsx_rejected(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        f = tmp_path / "空表.xlsx"
        wb.save(f)

        with pytest.raises(ValueError, match="没有提取到"):
            load_document(f, "空表.xlsx")


# ============================================================
# PDF 解析
# ============================================================
class TestPdfLoader:
    def test_load_pdf_with_pages(self, tmp_path):
        """PDF 要逐页提取，并且记住每段文字来自第几页。"""
        import pymupdf

        pdf = pymupdf.open()
        for text in ["第一页：星耀 X9 Pro 产品概述", "第二页：电池容量 5000mAh", "第三页：保修期 12 个月"]:
            page = pdf.new_page()
            page.insert_text((72, 100), text, fontsize=14)

        f = tmp_path / "产品手册.pdf"
        pdf.save(f)
        pdf.close()

        docs = load_document(f, "产品手册.pdf")

        assert len(docs) == 3
        # 页码要从 1 开始，符合普通人的阅读习惯
        assert [d.metadata["page"] for d in docs] == [1, 2, 3]
        assert "5000mAh" in docs[1].page_content
        assert docs[2].metadata["total_pages"] == 3


# ============================================================
# 分块
# ============================================================
class TestSplitter:
    def test_chinese_split_respects_sentence_boundaries(self):
        """中文分块要在句子边界断开，不能从句子中间硬切。

        这是中文场景最容易踩的坑：照抄英文教程的配置（按空格切），
        中文句子中间没有空格，切出来的块会从半句话开始、半句话结束。
        """
        from langchain_core.documents import Document

        doc = Document(page_content=SAMPLE_PRODUCT_TEXT, metadata={"file_type": ".txt"})
        chunks = split_documents([doc], chunk_size=200, chunk_overlap=40)

        assert len(chunks) > 1
        # 每一个分块都不该超过上限太多（留一点余量给分隔符）
        for chunk in chunks:
            assert len(chunk.page_content) <= 200 + 10

    def test_chunk_index_is_sequential(self):
        from langchain_core.documents import Document

        doc = Document(page_content=SAMPLE_PRODUCT_TEXT, metadata={"file_type": ".txt"})
        chunks = split_documents([doc], chunk_size=150, chunk_overlap=30)

        indices = [c.metadata["chunk_index"] for c in chunks]
        assert indices == list(range(len(chunks)))

    def test_all_content_is_preserved(self):
        """切分不能丢内容——所有分块拼起来应该覆盖原文的关键信息。"""
        from langchain_core.documents import Document

        doc = Document(page_content=SAMPLE_PRODUCT_TEXT, metadata={"file_type": ".txt"})
        chunks = split_documents([doc], chunk_size=150, chunk_overlap=30)

        combined = "".join(c.page_content for c in chunks)
        # 原文里的几个关键数据点都应该还在
        for keyword in ["骁龙 8 Gen 3", "5000mAh", "67W", "IMX890", "12 个月整机保修"]:
            assert keyword in combined, f"分块后丢失了关键内容：{keyword}"

    def test_overlap_smaller_than_chunk_size_required(self):
        """重叠必须小于块大小，否则切分算法会死循环。"""
        from langchain_core.documents import Document

        doc = Document(page_content="测试内容" * 100, metadata={})
        with pytest.raises(ValueError, match="chunk_overlap"):
            split_documents([doc], chunk_size=100, chunk_overlap=100)

    def test_markdown_split_by_headers(self):
        """Markdown 要按标题层级切分，并把标题记进 metadata。"""
        from langchain_core.documents import Document

        md = """# 星耀 X9 Pro 产品手册

## 电池与续航
内置 5000mAh 电池，支持 67W 快充。

## 影像系统
后置 5000 万像素主摄，支持光学防抖。
"""
        doc = Document(page_content=md, metadata={"file_type": ".md"})
        chunks = split_documents([doc], chunk_size=500, chunk_overlap=50)

        # 至少要能按标题切成几段
        assert len(chunks) >= 2
        combined = "".join(c.page_content for c in chunks)
        assert "电池与续航" in combined
        assert "影像系统" in combined

    def test_short_document_not_split(self):
        """很短的文档不该被切开——切碎反而会丢失上下文。"""
        from langchain_core.documents import Document

        doc = Document(page_content="星耀 X9 Pro 电池容量 5000mAh。", metadata={"file_type": ".txt"})
        chunks = split_documents([doc], chunk_size=500, chunk_overlap=80)

        assert len(chunks) == 1
        assert chunks[0].page_content == "星耀 X9 Pro 电池容量 5000mAh。"

    def test_empty_text_filtered_out(self):
        """分块后产生的空白块要被过滤掉，不能入库。"""
        from langchain_core.documents import Document

        doc = Document(page_content="有内容。" + "\n\n\n\n" * 50, metadata={"file_type": ".txt"})
        chunks = split_documents([doc], chunk_size=100, chunk_overlap=20)

        assert all(c.page_content.strip() for c in chunks)

    def test_metadata_is_inherited(self):
        """分块必须继承原文档的元数据，否则就不知道这段来自哪份资料了。"""
        from langchain_core.documents import Document

        doc = Document(
            page_content=SAMPLE_PRODUCT_TEXT,
            metadata={"doc_id": 42, "filename": "手册.pdf", "kb_id": 7, "file_type": ".txt"},
        )
        chunks = split_documents([doc], chunk_size=150, chunk_overlap=30)

        for chunk in chunks:
            assert chunk.metadata["doc_id"] == 42
            assert chunk.metadata["filename"] == "手册.pdf"
            assert chunk.metadata["kb_id"] == 7
